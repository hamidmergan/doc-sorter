"""Read the text of a document. Scans and images go through doc-extract."""

import tempfile
from pathlib import Path
from typing import List, Tuple

from .config import Settings
from .errors import LimitError, ModelError, UnsupportedFileError

TEXT_SUFFIXES = {".txt", ".md", ".markdown"}
DOCX_SUFFIX = ".docx"
PDF_SUFFIX = ".pdf"
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}

# a PDF with less text than this is treated as a scan
MIN_TEXT_CHARS = 30

VISION_HINT = " Scans and images need a vision-capable model (--vision-model)."


def supported_suffixes() -> List[str]:
    return sorted(TEXT_SUFFIXES | IMAGE_SUFFIXES | {DOCX_SUFFIX, PDF_SUFFIX})


def is_supported(path: Path) -> bool:
    return path.suffix.lower() in supported_suffixes()


def read_text(path: Path, settings: Settings, client=None) -> Tuple[str, str]:
    """Return (text, source), source is "text" or "vision". The text is cut to max_chars."""
    if not path.is_file():
        raise UnsupportedFileError(f"File not found: {path}")
    if not is_supported(path):
        raise UnsupportedFileError(f"File type {path.suffix or '(no extension)'} is not supported.")
    size_mb = path.stat().st_size / (1024 * 1024)
    if size_mb > settings.max_file_mb:
        raise LimitError(
            f"{path.name} is {size_mb:.1f} MB. The limit is {settings.max_file_mb:g} MB (max_file_mb)."
        )

    suffix = path.suffix.lower()
    source = "text"
    if suffix in TEXT_SUFFIXES:
        text = path.read_text(encoding="utf-8", errors="replace")
    elif suffix == DOCX_SUFFIX:
        text = _docx_text(path)
    elif suffix == PDF_SUFFIX:
        text = _pdf_text(path, settings.scan_pages)
        if len("".join(text.split())) < MIN_TEXT_CHARS:
            text, source = _scan_pdf(path, settings, client), "vision"
    else:
        text, source = _transcribe(path, settings, client), "vision"
    return text.strip()[: settings.max_chars], source


def _docx_text(path: Path) -> str:
    import docx

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise UnsupportedFileError(f"{path.name} cannot be read as a DOCX file: {exc}")
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(line for line in lines if line.strip())


def _open_pdf(path: Path):
    import pypdfium2 as pdfium

    try:
        return pdfium.PdfDocument(str(path))
    except Exception as exc:
        raise UnsupportedFileError(f"{path.name} cannot be read as a PDF: {exc}")


def _pdf_text(path: Path, pages: int) -> str:
    pdf = _open_pdf(path)
    try:
        texts = []
        for index in range(min(pages, len(pdf))):
            page = pdf[index]
            textpage = page.get_textpage()
            try:
                texts.append(textpage.get_text_range())
            finally:
                textpage.close()
                page.close()
        return "\n\n".join(texts)
    finally:
        pdf.close()


def _scan_pdf(path: Path, settings: Settings, client) -> str:
    """Transcribe only the first scan_pages pages; that is enough to sort a document."""
    import pypdfium2 as pdfium

    pdf = _open_pdf(path)
    try:
        if len(pdf) <= settings.scan_pages:
            return _transcribe(path, settings, client)
        head = pdfium.PdfDocument.new()
        try:
            head.import_pages(pdf, list(range(settings.scan_pages)))
            with tempfile.TemporaryDirectory() as tmp:
                part = Path(tmp) / path.name
                head.save(str(part))
                return _transcribe(part, settings, client)
        finally:
            head.close()
    finally:
        pdf.close()


def _transcribe(path: Path, settings: Settings, client) -> str:
    import doc_extract

    try:
        return doc_extract.extract(
            path,
            base_url=settings.base_url,
            model=settings.vision_model,
            api_key=settings.api_key,
            client=client,
            timeout=settings.timeout,
            max_retries=settings.max_retries,
            max_pages=settings.scan_pages,
            max_file_mb=settings.max_file_mb,
        )
    except doc_extract.ModelError as exc:
        raise ModelError(str(exc) + VISION_HINT)
    except doc_extract.ExtractError as exc:
        raise UnsupportedFileError(str(exc))
